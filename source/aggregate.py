import csv
import os
import sys
from datetime import datetime, timedelta

if getattr(sys, "frozen", False):
    # .exe化されている場合: exe自身が置かれている場所を基準にする
    BASE_DIR = os.path.dirname(sys.executable)
else:
    # 開発時: このファイルは 納品物/source/ にあるので、1つ上(納品物/)を基準にする
    BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

DATA_DIR = os.path.join(BASE_DIR, "data")
OUTPUT_DIR = os.path.join(BASE_DIR, "output")

STORES = [
    {"label": "店A", "filename": "店A_売上.csv", "encoding": "utf-8", "closed_weekdays": [2]},  # 2=水曜
    {"label": "店B", "filename": "店B_売上.csv", "encoding": "utf-8", "closed_weekdays": []},
    {"label": "店C", "filename": "店C_売上.csv", "encoding": "cp932", "closed_weekdays": []},
]

CLOSED_WEEKDAYS = {store["label"]: store["closed_weekdays"] for store in STORES}


DATE_FORMATS = ["%Y-%m-%d", "%Y/%m/%d"]


def normalize_date(raw):
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(raw.strip(), fmt).date()
        except ValueError:
            continue
    raise ValueError(f"日付の形式が想定外です: {raw!r}")


def clean_amount(raw):
    cleaned = raw.strip().strip("¥").replace(",", "").replace("¥", "")
    return int(cleaned)


def read_store(store):
    """日付(date) -> {"来客数": int, "売上金額": int} の辞書を返す"""
    data = {}
    path = os.path.join(DATA_DIR, store["filename"])
    with open(path, encoding=store["encoding"], newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            date = normalize_date(row["日付"])
            data[date] = {
                "来客数": int(row["来客数"]),
                "売上金額": clean_amount(row["売上金額"]),
            }
    return data


def target_date(stores_data):
    """全店舗のデータの中で最も新しい日付を対象日とする"""
    return max(date for data in stores_data.values() for date in data)


def previous_available_date(store_data, before_date):
    """store_dataの中で、before_dateより前で最も新しい日付を返す（無ければNone）"""
    candidates = [d for d in store_data if d < before_date]
    return max(candidates) if candidates else None


def compute_summary(stores_data, target):
    """店ごとの対象日・直前営業日の売上と、全体合計・差額・増減率をまとめる"""
    per_store = {}
    for label, data in stores_data.items():
        baseline_date = previous_available_date(data, target)
        absent = target not in data
        is_known_closure = absent and target.weekday() in CLOSED_WEEKDAYS.get(label, [])
        per_store[label] = {
            "対象日来客数": data.get(target, {}).get("来客数"),
            "対象日売上": data.get(target, {}).get("売上金額"),
            # 既知の定休日(曜日固定)による欠落と、それ以外の原因不明の欠落(データ未着など)を区別する。
            # 区別せずに一律「定休日」にすると、送り忘れのような異常を見逃す。
            "状態": "定休日" if is_known_closure else ("データ未着" if absent else None),
            "直前営業日": baseline_date,
            "直前営業日売上": data.get(baseline_date, {}).get("売上金額") if baseline_date else None,
        }

    # 対象日にデータが無い店(定休日)は、比較の両側から除外する。
    # 片側だけに含めると「休業しただけ」が「急落」に見えてしまう。
    active = [v for v in per_store.values() if v["対象日売上"] is not None]
    total_target = sum(v["対象日売上"] for v in active)
    total_baseline = sum(v["直前営業日売上"] for v in active if v["直前営業日売上"] is not None)
    diff = total_target - total_baseline
    pct = (diff / total_baseline * 100) if total_baseline else None

    return {
        "対象日": target,
        "店舗別": per_store,
        "合計_対象日": total_target,
        "合計_直前営業日": total_baseline,
        "差額": diff,
        "増減率": pct,
    }


def write_report(summary, output_dir=OUTPUT_DIR):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    target = summary["対象日"]
    wb = Workbook()
    ws = wb.active
    ws.title = "日次売上サマリ"

    ws["A1"] = f"日次売上サマリ（対象日: {target}）"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"作成日時: {datetime.now().strftime('%Y-%m-%d %H:%M')}"

    header_row = 4
    for col, text in enumerate(["店舗", "来客数", "売上金額", "備考"], start=1):
        ws.cell(row=header_row, column=col, value=text).font = Font(bold=True)

    row = header_row + 1
    note_lines = []
    warning_lines = []
    for label, info in summary["店舗別"].items():
        ws.cell(row=row, column=1, value=label)
        if info["状態"]:
            ws.cell(row=row, column=2, value="—")
            ws.cell(row=row, column=3, value="—")
            ws.cell(row=row, column=4, value=info["状態"])
            if info["状態"] == "データ未着":
                warning_lines.append(f"⚠ {label}: 本日分のデータが届いていません。ご確認ください。")
        else:
            ws.cell(row=row, column=2, value=info["対象日来客数"])
            ws.cell(row=row, column=3, value=info["対象日売上"]).number_format = "¥#,##0"

        prev = info["直前営業日"]
        if prev is not None and prev != target - timedelta(days=1):
            note_lines.append(f"{label}: 直前の営業日は{prev}（間に定休日を含む）")
        row += 1

    total_row = row
    ws.cell(row=total_row, column=1, value="合計").font = Font(bold=True)
    total_count = sum(
        info["対象日来客数"] for info in summary["店舗別"].values() if info["対象日売上"] is not None
    )
    ws.cell(row=total_row, column=2, value=total_count).font = Font(bold=True)
    total_cell = ws.cell(row=total_row, column=3, value=summary["合計_対象日"])
    total_cell.font = Font(bold=True)
    total_cell.number_format = "¥#,##0"

    ws.cell(row=total_row + 2, column=1, value="前営業日合計")
    ws.cell(row=total_row + 2, column=3, value=summary["合計_直前営業日"]).number_format = "¥#,##0"

    ws.cell(row=total_row + 3, column=1, value="差額")
    ws.cell(row=total_row + 3, column=3, value=summary["差額"]).number_format = "¥#,##0;[Red]-¥#,##0"

    ws.cell(row=total_row + 4, column=1, value="増減率")
    pct = summary["増減率"] / 100 if summary["増減率"] is not None else None
    ws.cell(row=total_row + 4, column=3, value=pct).number_format = "0.0%"

    footer_row = total_row + 6
    for i, line in enumerate(warning_lines):
        cell = ws.cell(row=footer_row + i, column=1, value=line)
        cell.font = Font(bold=True, color="CC0000")
    footer_row += len(warning_lines) + (1 if warning_lines else 0)

    for i, line in enumerate(note_lines):
        ws.cell(row=footer_row + i, column=1, value=line)

    for col, width in zip("ABCD", [10, 10, 14, 32]):
        ws.column_dimensions[col].width = width

    os.makedirs(output_dir, exist_ok=True)
    path = os.path.join(output_dir, f"日次売上サマリ_{target.isoformat()}.xlsx")
    wb.save(path)
    return path


if __name__ == "__main__":
    try:
        stores_data = {store["label"]: read_store(store) for store in STORES}
        target = target_date(stores_data)
        summary = compute_summary(stores_data, target)

        path = write_report(summary)
        print(f"レポートを出力しました: {path}")
    except Exception as e:
        print("エラーが発生し、処理を中断しました。")
        print(f"詳細: {e}")
        print("『data』フォルダの中のファイル名・中身をご確認のうえ、直らない場合はご連絡ください。")
        input("Enterキーを押すとこの画面を閉じます...")
        sys.exit(1)
